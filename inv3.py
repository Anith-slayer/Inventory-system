import sys
import os
import openpyxl
from openpyxl import load_workbook
from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, 
    QTableWidget, QTableWidgetItem, QPushButton, QLineEdit, 
    QLabel, QMessageBox, QDialog, QFormLayout, QComboBox
)
from PyQt6.QtGui import QFont, QPalette, QColor
from PyQt6.QtCore import Qt

class InventoryManagementSystem(QMainWindow):
    def __init__(self):
        super().__init__()
        self.file_name = "inventory.xlsx"
        self.init_workbook()
        self.init_ui()

    def init_workbook(self):
        """Initialize the Excel workbook if it doesn't exist."""
        if not os.path.exists(self.file_name):
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Inventory"
            sheet.append(["Category", "Subcategory", "Bike Model", "Product Name", "Price", "Quantity"])
            workbook.save(self.file_name)

    def init_ui(self):
        """Set up the main user interface."""
        self.setWindowTitle("Bike Inventory Management System")
        self.setGeometry(100, 100, 1000, 600)
        
        # Central widget and main layout
        central_widget = QWidget()
        main_layout = QVBoxLayout()
        
        # Title
        title_label = QLabel("Bike Inventory Management System")
        title_label.setFont(QFont('Arial', 16, QFont.Weight.Bold))
        title_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        main_layout.addWidget(title_label)
        
        # Search and Filter Layout
        search_layout = QHBoxLayout()
        
        # Search Input
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("Search Product Name")
        search_layout.addWidget(self.search_input)
        
        # Category Filter
        self.category_filter = QComboBox()
        self.category_filter.addItem("All Categories")
        search_layout.addWidget(self.category_filter)
        
        # Search Button
        search_button = QPushButton("Search")
        search_button.clicked.connect(self.search_products)
        search_layout.addWidget(search_button)
        
        main_layout.addLayout(search_layout)
        
        # Table for displaying inventory
        self.table = QTableWidget()
        self.table.setColumnCount(6)
        self.table.setHorizontalHeaderLabels(["Category", "Subcategory", "Bike Model", "Product Name", "Price", "Quantity"])
        self.table.horizontalHeader().setStretchLastSection(True)
        main_layout.addWidget(self.table)
        
        # Button Layout
        button_layout = QHBoxLayout()
        buttons = [
            ("Add Product", self.open_add_product_dialog),
            ("Increase Stock", self.open_stock_dialog),
            ("Reduce Stock", self.open_stock_dialog),
            ("Refresh", self.load_inventory)
        ]
        
        for label, method in buttons:
            btn = QPushButton(label)
            btn.clicked.connect(method)
            button_layout.addWidget(btn)
        
        main_layout.addLayout(button_layout)
        
        central_widget.setLayout(main_layout)
        self.setCentralWidget(central_widget)
        
        # Load initial inventory
        self.load_inventory()
        self.populate_categories()

    def load_inventory(self):
        """Load inventory data into the table."""
        workbook = load_workbook(self.file_name)
        sheet = workbook["Inventory"]
        
        # Clear existing rows
        self.table.setRowCount(0)
        
        # Populate table
        for row in sheet.iter_rows(min_row=2, values_only=True):
            rowPosition = self.table.rowCount()
            self.table.insertRow(rowPosition)
            for i, value in enumerate(row):
                self.table.setItem(rowPosition, i, QTableWidgetItem(str(value)))
        
        workbook.close()

    def populate_categories(self):
        """Populate the category filter dropdown."""
        categories = set()
        workbook = load_workbook(self.file_name)
        sheet = workbook["Inventory"]
        
        for row in sheet.iter_rows(min_row=2, values_only=True):
            categories.add(str(row[0]))
        
        workbook.close()
        
        self.category_filter.clear()
        self.category_filter.addItem("All Categories")
        for category in sorted(categories):
            self.category_filter.addItem(category)

    def search_products(self):
        """Search and filter products based on user input."""
        search_term = self.search_input.text().lower()
        selected_category = self.category_filter.currentText()
        
        for row in range(self.table.rowCount()):
            category = self.table.item(row, 0).text()
            product_name = self.table.item(row, 3).text()
            
            # Apply category and search term filters
            category_match = selected_category == "All Categories" or category == selected_category
            name_match = search_term == "" or search_term in product_name.lower()
            
            self.table.setRowHidden(row, not (category_match and name_match))

    def open_add_product_dialog(self):
        """Open a dialog to add a new product."""
        dialog = QDialog(self)
        dialog.setWindowTitle("Add Product")
        form_layout = QFormLayout()
        
        # Input fields for product details
        category_input = QLineEdit()
        subcategory_input = QLineEdit()
        bike_model_input = QLineEdit()
        product_name_input = QLineEdit()
        price_input = QLineEdit()
        quantity_input = QLineEdit()
        
        form_layout.addRow("Category:", category_input)
        form_layout.addRow("Subcategory:", subcategory_input)
        form_layout.addRow("Bike Model:", bike_model_input)
        form_layout.addRow("Product Name:", product_name_input)
        form_layout.addRow("Price:", price_input)
        form_layout.addRow("Quantity:", quantity_input)
        
        # Buttons for submission
        button_layout = QHBoxLayout()
        submit_button = QPushButton("Add")
        cancel_button = QPushButton("Cancel")
        button_layout.addWidget(submit_button)
        button_layout.addWidget(cancel_button)
        form_layout.addRow(button_layout)
        
        dialog.setLayout(form_layout)
        
        def add_product():
            """Add the product to the inventory file."""
            category = category_input.text()
            subcategory = subcategory_input.text()
            bike_model = bike_model_input.text()
            product_name = product_name_input.text()
            try:
                price = float(price_input.text())
                quantity = int(quantity_input.text())
            except ValueError:
                QMessageBox.warning(self, "Input Error", "Price must be a number and Quantity must be an integer.")
                return
            
            if not all([category, subcategory, bike_model, product_name]):
                QMessageBox.warning(self, "Input Error", "All fields are required.")
                return
            
            # Append to Excel
            workbook = load_workbook(self.file_name)
            sheet = workbook["Inventory"]
            sheet.append([category, subcategory, bike_model, product_name, price, quantity])
            workbook.save(self.file_name)
            workbook.close()
            
            QMessageBox.information(self, "Success", "Product added successfully!")
            dialog.accept()
            self.load_inventory()
        
        submit_button.clicked.connect(add_product)
        cancel_button.clicked.connect(dialog.reject)
        
        dialog.exec()

    def open_stock_dialog(self):
        """Placeholder for stock dialog functionality."""
        QMessageBox.information(self, "Info", "This functionality is under development.")

def main():
    app = QApplication(sys.argv)
    
    # Optional: Set a dark theme
    app.setStyle('Fusion')
    dark_palette = QPalette()
    dark_palette.setColor(QPalette.ColorRole.Window, QColor(53, 53, 53))
    dark_palette.setColor(QPalette.ColorRole.WindowText, Qt.GlobalColor.white)
    dark_palette.setColor(QPalette.ColorRole.Base, QColor(25, 25, 25))
    dark_palette.setColor(QPalette.ColorRole.AlternateBase, QColor(53, 53, 53))
    dark_palette.setColor(QPalette.ColorRole.ToolTipBase, Qt.GlobalColor.white)
    dark_palette.setColor(QPalette.ColorRole.ToolTipText, Qt.GlobalColor.white)
    dark_palette.setColor(QPalette.ColorRole.Text, Qt.GlobalColor.white)
    dark_palette.setColor(QPalette.ColorRole.Button, QColor(53, 53, 53))
    dark_palette.setColor(QPalette.ColorRole.ButtonText, Qt.GlobalColor.white)
    dark_palette.setColor(QPalette.ColorRole.BrightText, Qt.GlobalColor.red)
    dark_palette.setColor(QPalette.ColorRole.Link, QColor(42, 130, 218))
    dark_palette.setColor(QPalette.ColorRole.Highlight, QColor(42, 130, 218))
    dark_palette.setColor(QPalette.ColorRole.HighlightedText, Qt.GlobalColor.black)
    app.setPalette(dark_palette)
    
    inventory_system = InventoryManagementSystem()
    inventory_system.show()
    sys.exit(app.exec())

if __name__ == "__main__":
    main()

